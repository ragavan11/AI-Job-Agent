"""DOCX export via python-docx — same content model as pdf_export.py,
kept in a separate module since the two libraries share no code."""
from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt, RGBColor

from app.models import Profile, Resume

_INK = RGBColor(0x14, 0x18, 0x1F)
_MUTED = RGBColor(0x5B, 0x64, 0x72)
_ACCENT = RGBColor(0x1F, 0x6F, 0x78)


def resume_to_docx(profile: Profile, resume: Resume) -> bytes:
    doc = Document()

    name = doc.add_paragraph()
    run = name.add_run(profile.full_name)
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = _INK

    contact = doc.add_paragraph()
    contact_bits = [b for b in [profile.email, profile.phone, profile.linkedin_url, profile.github_url, profile.portfolio_url] if b]
    run = contact.add_run(" | ".join(contact_bits))
    run.font.size = Pt(9.5)
    run.font.color.rgb = _MUTED

    def section(title: str):
        p = doc.add_paragraph()
        r = p.add_run(title.upper())
        r.font.bold = True
        r.font.size = Pt(11.5)
        r.font.color.rgb = _ACCENT

    if resume.summary:
        section("Summary")
        doc.add_paragraph(resume.summary)

    if resume.skills:
        section("Skills")
        doc.add_paragraph(" • ".join(resume.skills))

    if resume.work_history:
        section("Experience")
        for job in resume.work_history:
            p = doc.add_paragraph()
            r = p.add_run(f"{job.get('title', '')} — {job.get('company', '')}")
            r.font.bold = True
            meta = doc.add_paragraph()
            meta_run = meta.add_run(f"{job.get('start', '')} – {job.get('end', 'Present')}")
            meta_run.font.size = Pt(9)
            meta_run.font.color.rgb = _MUTED
            for bullet in job.get("bullets", []):
                doc.add_paragraph(bullet, style="List Bullet")

    if resume.projects:
        section("Projects")
        for proj in resume.projects:
            p = doc.add_paragraph()
            tech = ", ".join(proj.get("tech", []))
            r = p.add_run(f"{proj.get('name', '')}" + (f" ({tech})" if tech else ""))
            r.font.bold = True
            doc.add_paragraph(proj.get("description", ""))

    if resume.education:
        section("Education")
        for edu in resume.education:
            doc.add_paragraph(f"{edu.get('degree', '')} — {edu.get('school', '')} ({edu.get('year', '')})")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
